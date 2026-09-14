# -*- coding: utf-8 -*-
"""Supplementary material for the paper, which carries no tables or figures of its own. Writes supplementary/:
Table1.md, Table2.md, Table3.md (each with its caption and a CSV twin), Figure1.png and Figure1.pdf (copied from
results/main-1/paper/), a README index, and Supplementary_Material.docx, which Word exports to PDF.

    python scripts/supplementary.py        (run paper_figure_single.py and keyword_screen.py first)
"""
import csv
import shutil

from config import RESULTS, ROOT
from tables import write_csv

OUT = ROOT / "supplementary"
REPO = "https://github.com/ymaou/language-legal-outcome-variation/blob/main"
LABEL = {"terra": "GPT-5.6 Terra", "terra-off": "GPT-5.6 Terra, reasoning off", "sonnet": "Claude Sonnet 5",
         "gemini": "Gemini 3.6 Flash"}


def rd(run, name):
    with open(RESULTS / run / "csv" / f"{name}.csv", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def n2(x):
    s = f"{float(x):.2f}"
    return "0.00" if s == "-0.00" else s.replace("-", "−")


def ci(value, low, high):
    return f"{n2(value)} [{n2(low)}, {n2(high)}]"


def greek_counts(run, cfg):
    bs = [b for b in rd(run, "blocks") if b["configuration"] == cfg and b["arm"] == "gr"]
    n = sum(int(b["valid draws used"]) for b in bs)
    return n, sum(int(b["option 3"]) for b in bs), sum(int(b["option 1"]) + int(b["option 2"]) for b in bs)


def primary_rates(run, cfg):
    return {r["measure"]: r for r in rd(run, "rates") if r["configuration"] == cfg and r["definition"].startswith("primary")}


def table1():
    headers = ["Model", "Greek draws dismissed", "Greek draws granting relief", "Agreement with court", "MCC", "F",
               "W_en", "L_EN", "Δ_EN [95% CI]", "W_de", "L_DE", "Mean D [95% CI]", "Ground differs: EN / BT-en"]
    rows = []
    for cfg, name in LABEL.items():
        n, dis, grant = greek_counts("main-1", cfg)
        s = next(r for r in rd("main-1", "sensitivity") if r["configuration"] == cfg)
        r = primary_rates("main-1", cfg)
        d = next(x for x in rd("main-1", "direction") if x["configuration"] == cfg)
        g = next(x for x in rd("main-1", "ground") if x["configuration"] == cfg)
        rows.append([name, f"{dis}/{n}", f"{grant}/{n}", n2(s["accuracy (four options)"]), n2(s["MCC"]),
                     n2(r["F"]["value"]), n2(r["W_en"]["value"]), n2(r["L_EN"]["value"]),
                     ci(r["delta_EN"]["value"], r["delta_EN"]["95% CI low"], r["delta_EN"]["95% CI high"]),
                     n2(r["W_de"]["value"]), n2(r["L_DE"]["value"]),
                     ci(d["mean D (EN grant share minus Greek)"], d["95% CI low"], d["95% CI high"]),
                     f"{n2(g['share (EN)'])} / {n2(g['share (BT-en)'])}"])
    caption = ("**Supplementary Table 1.** Main run (the model decides the application as a judge): what each model "
               "decided in Greek, its agreement with the courts' orders, and the pre-specified rates. Each model decided "
               "each case 31 times in each version and 62 times in Greek. Agreement with court is the share of the ten "
               "cases on which the Greek modal outcome matched the court's order (majority-class baseline 0.40); MCC is "
               "the Matthews correlation coefficient. F (noise floor), W (rewording rate) and L (language rate) are "
               "shares of the ten cases on which the modal outcome changed: between two random halves of the Greek "
               "draws (F), between a Greek half and the back-translation (W), and between a Greek half and the "
               "translation (L); \"en\" and \"de\" denote the English and German routes. Δ_EN = L_EN − W_en, with a "
               "95% interval from a paired bootstrap over cases (10,000 resamples). Mean D is the share of draws "
               "granting relief in English minus that share across all Greek draws, averaged over cases. \"Ground "
               "differs\" is the share of comparisons, among those in which the modal outcome held, in which the modal "
               "ground differed.")
    return "Table1", caption, headers, rows


def table2():
    headers = ["Model", "Framing", "Greek draws dismissed", "Agreement with court", "F", "L_EN", "Mean D [95% CI]",
               "Cases flipping to refusal in English"]
    rows = []
    for cfg, name in LABEL.items():
        for run, framing in (("main-1", "decide"), ("predict-1", "predict")):
            n, dis, _ = greek_counts(run, cfg)
            s = next(r for r in rd(run, "sensitivity") if r["configuration"] == cfg)
            r = primary_rates(run, cfg)
            d = next(x for x in rd(run, "direction") if x["configuration"] == cfg)
            rows.append([name, framing, f"{dis}/{n}", n2(s["accuracy (four options)"]), n2(r["F"]["value"]),
                         n2(r["L_EN"]["value"]), ci(d["mean D (EN grant share minus Greek)"], d["95% CI low"], d["95% CI high"]),
                         d["flips towards refusing"]])
    caption = ("**Supplementary Table 2.** Decide against predict, Greek and English versions. Under \"decide\" (main "
               "run) the model decides the application as a judge, 31 draws per version and 62 in Greek; under "
               "\"predict\" (exploratory run, designed after the main results were known) it predicts the court's "
               "decision as a legal analyst, 15 draws per version and 30 in Greek. Measures as in Supplementary "
               "Table 1. The predictive run had no back-translation version, so its language rate is not net of "
               "rewording.")
    return "Table2", caption, headers, rows


def table3():
    with open(RESULTS / "main-1" / "paper" / "T3_considerations.csv", encoding="utf-8-sig") as f:
        t3 = list(csv.DictReader(f))
    considerations = list(dict.fromkeys(r["consideration"] for r in t3))
    headers = ["Consideration invoked in the reasoning"] + list(LABEL.values())
    rows = []
    for c in considerations:
        cells = []
        for cfg in LABEL:
            r = next(x for x in t3 if x["configuration"] == cfg and x["consideration"] == c)
            cells.append(f"{float(r['share of all replies']):.0%}")
        rows.append([c[0].upper() + c[1:]] + cells)
    caption = ("**Supplementary Table 3.** Considerations invoked in the models' reasoning in the main run: the share "
               "of each model's 2,170 valid replies, across all ten cases and six versions, in which an automated "
               "keyword screen in Greek, English and German matched the consideration. A screen, not hand coding; the "
               f"patterns are in [scripts/keyword_screen.py]({REPO}/scripts/keyword_screen.py).")
    return "Table3", caption, headers, rows


FIGURE_CAPTION = ("**Supplementary Figure 1.** Modal outcome per case for each model, with the court's order. Each "
                  "cell is the outcome chosen most often across the draws for that case: 31 per version (62 for "
                  "Greek) under the decide framing, 15 (30 for Greek) under the predict framing. GR: Greek original; "
                  "BT-en: Greek back-translation of the English; EN: English; BT-de: Greek back-translation of the "
                  "German; DE-lit: German with German Part 24 terms; DE-eng: German with the English Part 24 terms.")


def markdown_table(headers, rows):
    return "\n".join(["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
                     + ["| " + " | ".join(str(x) for x in r) + " |" for r in rows])


def main():
    OUT.mkdir(exist_ok=True)
    tables = [table1(), table2(), table3()]
    for stem, caption, headers, rows in tables:
        (OUT / f"{stem}.md").write_text(caption + "\n\n" + markdown_table(headers, rows) + "\n", encoding="utf-8")
        write_csv(OUT / f"{stem}.csv", headers, rows)
    for ext in ("png", "pdf"):
        shutil.copyfile(RESULTS / "main-1" / "paper" / f"Figure1.{ext}", OUT / f"Figure1.{ext}")
    (OUT / "Figure1.md").write_text(f"{FIGURE_CAPTION}\n\n![Supplementary Figure 1](Figure1.png)\n", encoding="utf-8")

    index = ["# Supplementary material",
             "",
             "*Does language change the outcome? A trilingual study of language-conditioned outcome variation in LLM "
             "legal reasoning.* The paper carries no tables or figures; they are here. Every number is computed by "
             "script from the released model outputs (`raw/`); see `results/README.md` for how to regenerate each one.",
             "",
             "| Item | File | Also as |",
             "|---|---|---|",
             "| Supplementary Table 1: main run, per model | [Table1.md](Table1.md) | [CSV](Table1.csv) |",
             "| Supplementary Table 2: decide against predict | [Table2.md](Table2.md) | [CSV](Table2.csv) |",
             "| Supplementary Table 3: considerations in the reasoning | [Table3.md](Table3.md) | [CSV](Table3.csv) |",
             "| Supplementary Figure 1: modal outcome per case, both framings | [Figure1.md](Figure1.md) | [PNG](Figure1.png), [PDF](Figure1.pdf) |",
             "| All of the above in one document | [Supplementary_Material.pdf](Supplementary_Material.pdf) | [DOCX](Supplementary_Material.docx) |",
             "",
             "Further comparisons (per case, per version, inter-model agreement, grounds, reliability, tokens and cost) "
             "are in `results/main-1/package/` and `results/predict-1/package/`."]
    (OUT / "README.md").write_text("\n".join(index) + "\n", encoding="utf-8")

    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Cm, Pt
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(9)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.6))
    title = doc.add_paragraph()
    run = title.add_run("Supplementary material. Does language change the outcome? A trilingual study of "
                        "language-conditioned outcome variation in LLM legal reasoning")
    run.bold = True
    doc.add_paragraph("Yiolanti Maou. Repository: https://github.com/ymaou/language-legal-outcome-variation")
    for stem, caption, headers, rows in tables:
        p = doc.add_paragraph()
        head, _, rest = caption.partition("** ")
        p.add_run(head.strip("*") + " ").bold = True
        p.add_run(rest.replace("[scripts/keyword_screen.py](" + REPO + "/scripts/keyword_screen.py)",
                               "scripts/keyword_screen.py"))
        tb = doc.add_table(rows=1, cols=len(headers))
        tb.style = "Table Grid"
        for i, h in enumerate(headers):
            cell = tb.rows[0].cells[i]
            cell.text = h
            for r in cell.paragraphs[0].runs:
                r.bold = True
                r.font.size = Pt(8)
        for row in rows:
            cells = tb.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = str(v)
                for r in cells[i].paragraphs[0].runs:
                    r.font.size = Pt(8)
        doc.add_paragraph()
    doc.add_page_break()
    head, _, rest = FIGURE_CAPTION.partition("** ")
    doc.add_picture(str(OUT / "Figure1.png"), width=Cm(15))
    p = doc.add_paragraph()
    p.add_run(head.strip("*") + " ").bold = True
    p.add_run(rest)
    doc.save(OUT / "Supplementary_Material.docx")
    print("written:", sorted(x.name for x in OUT.iterdir()))


if __name__ == "__main__":
    main()
