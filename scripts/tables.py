# -*- coding: utf-8 -*-
"""Write tables to an Excel workbook (one tab per table) and to CSV files."""
import csv
import json
import math

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CELL_LIMIT = 32000


def _value(v):
    if v is None:
        return None
    if isinstance(v, float):
        return None if math.isnan(v) else round(v, 10)
    if isinstance(v, (dict, list, tuple, set)):
        v = json.dumps(sorted(v) if isinstance(v, set) else v, ensure_ascii=False)
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        if len(v) > CELL_LIMIT:
            v = v[:CELL_LIMIT] + " [cut here: an Excel cell holds at most 32,767 characters]"
    return v


def write_workbook(path, sheets):
    """sheets: list of dicts with name, headers, rows; optional notes (lines above the table), widths
    ({header: width}), wrap (headers whose cells wrap) and validations ({header: "a,b,c"})."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(sheet["name"][:31])
        row = 1
        for note in sheet.get("notes", []):
            ws.cell(row=row, column=1, value=_value(note)).font = Font(italic=True)
            row += 1
        if sheet.get("notes"):
            row += 1
        header_row = row
        headers = sheet["headers"]
        for j, header in enumerate(headers, 1):
            ws.cell(row=header_row, column=j, value=header).font = Font(bold=True)
        wrap_columns = {headers.index(h) + 1 for h in sheet.get("wrap", []) if h in headers}
        for values in sheet["rows"]:
            row += 1
            for j, v in enumerate(values, 1):
                cell = ws.cell(row=row, column=j, value=_value(v))
                if j in wrap_columns:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
        widths = sheet.get("widths", {})
        for j, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(j)].width = widths.get(header, min(30, max(10, len(header) + 2)))
        last = max(row, header_row + 1)
        for header, options in sheet.get("validations", {}).items():
            if header not in headers:
                continue
            letter = get_column_letter(headers.index(header) + 1)
            dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}{header_row + 1}:{letter}{last}")
    wb.save(path)


def write_csv(path, headers, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for values in rows:
            w.writerow(["" if _value(v) is None else _value(v) for v in values])
